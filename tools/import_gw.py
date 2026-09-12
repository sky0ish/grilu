# -*- coding: utf-8 -*-
"""
그룹웨어(gw.gri.re.kr)에서 내려받은 data/gw/gw_export.zip (게시글 + 첨부파일) 을
1) 첨부파일 → files/gw/ 에 저장 (사이트에서 내려받기)
2) 첨부파일 텍스트 추출 (pdf: pymupdf, hwp: pyhwp, hwpx: xml, xls: xlrd)
3) data/gw_posts.json 으로 정리 (build.py 가 정적 페이지·목록 생성, seed SQL 생성)
실행: python tools/import_gw.py
"""
import os, re, io, json, html, zipfile, tempfile, sys
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ZIPS = sorted(__import__("glob").glob(os.path.join(ROOT, "data", "gw", "gw_export*.zip")))
FILES = os.path.join(ROOT, "files", "gw")
MAX_ATT = 95 * 1024 * 1024      # GitHub 파일 한도(100MB) 안의 첨부는 모두 저장 (그보다 큰 동영상만 그룹웨어 안내)
BOARD_MAP = {"000000371": ("committee", "심의위원회 상정(안)"), "0000001yl": ("rules", "노사협의회 운영규약"),
             "0000001ym": ("council", "공고 및 회의록"), "0000001yn": ("council", "안건 제안"), "0000001hy": ("director", "노동이사 활동보고"), "00000039t": ("budget", "예산결산서")}

def pdf_text(b):
    import fitz
    return "\n".join(p.get_text("text") for p in fitz.open(stream=b, filetype="pdf"))

def hwp_text(b):
    from hwp5.xmlmodel import Hwp5File
    from hwp5.hwp5txt import TextTransform
    tmp = tempfile.NamedTemporaryFile(suffix=".hwp", delete=False); tmp.write(b); tmp.close()
    out = io.BytesIO()
    hf = Hwp5File(tmp.name)
    try:
        TextTransform().transform_hwp5_to_text(hf, out)
    finally:
        try: hf.close()
        except Exception: pass
        try: os.unlink(tmp.name)
        except Exception: pass
    return out.getvalue().decode("utf-8", "replace")

def hwpx_text(b):
    zz = zipfile.ZipFile(io.BytesIO(b)); parts = []
    for n in sorted(zz.namelist()):
        if n.startswith("Contents/section"):
            x = zz.read(n).decode("utf-8", "replace")
            x = re.sub(r"</hp:p>", "\n", x); x = re.sub(r"<[^>]+>", "", x)
            parts.append(html.unescape(x))
    return "\n".join(parts)

def xlsx_text(b):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True, read_only=True); out = []
        for ws in wb.worksheets:
            out.append("[" + ws.title + "]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c not in (None, "")]
                if cells: out.append(chr(9).join(cells))
        return chr(10).join(out)
    except Exception as e:
        return ""

def xls_text(b):
    try:
        import xlrd
        wb = xlrd.open_workbook(file_contents=b); out = []
        for sh in wb.sheets():
            out.append("[" + sh.name + "]")
            for r in range(sh.nrows): out.append("\t".join(str(c) for c in sh.row_values(r) if c != ""))
        return "\n".join(out)
    except Exception as e:
        return ""

def extract(name, b):
    ext = os.path.splitext(name)[1].lower()
    try:
        if ext == ".pdf": return pdf_text(b)
        if ext == ".hwp": return hwp_text(b)
        if ext == ".hwpx": return hwpx_text(b)
        if ext == ".xls": return xls_text(b)
        if ext == ".xlsx": return xlsx_text(b)
    except Exception as e:
        print("   ! 텍스트 추출 실패", name, e)
    return ""

# ---------- 표를 살리는 HTML 변환 ----------
def _esc(t): return html.escape(t or "")

def pdf_html(b):
    """PDF → HTML: 표(find_tables)는 <table>, 나머지 글은 <p>. 페이지 순서·세로 위치 순으로 배치"""
    import fitz
    doc = fitz.open(stream=b, filetype="pdf"); out = []
    for page in doc:
        items = []
        try:
            tabs = page.find_tables()
        except Exception:
            tabs = None
        boxes = []
        if tabs:
            for t in tabs.tables:
                rows = t.extract()
                if not rows or len(rows) < 2: continue
                boxes.append(fitz.Rect(t.bbox))
                th = "".join(f"<th>{_esc(c).replace(chr(10), '<br>')}</th>" for c in rows[0])
                trs = "".join("<tr>" + "".join(f"<td>{_esc(c).replace(chr(10), '<br>')}</td>" for c in r) + "</tr>" for r in rows[1:])
                items.append((t.bbox[1], f"<table class='doc-table'><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>"))
        for blk in page.get_text("blocks"):
            x0, y0, x1, y1, txt = blk[0], blk[1], blk[2], blk[3], blk[4]
            r = fitz.Rect(x0, y0, x1, y1)
            if any(r.intersects(bx) and (r & bx).get_area() > r.get_area() * 0.5 for bx in boxes): continue
            txt = re.sub(r"[ \t]+", " ", txt).strip()
            if not txt or re.fullmatch(r"-?\s*\d{1,3}\s*-?", txt): continue
            items.append((y0, "".join(f"<p>{_esc(line)}</p>" for line in txt.split(chr(10)) if line.strip())))
        items.sort(key=lambda x: x[0])
        out.append("".join(h for _, h in items))
    return "<div class='doc-page'>" + "</div><div class='doc-page'>".join(out) + "</div>"

def hwp_html(b):
    """HWP → HTML (pyhwp hwp5html). 표·문단 구조 유지, 스타일은 제거"""
    from hwp5.xmlmodel import Hwp5File
    from hwp5.hwp5html import HTMLTransform
    tmp = tempfile.NamedTemporaryFile(suffix=".hwp", delete=False); tmp.write(b); tmp.close()
    d = tempfile.mkdtemp(); hf = Hwp5File(tmp.name)
    try:
        HTMLTransform().transform_hwp5_to_dir(hf, d)
        x = open(os.path.join(d, "index.xhtml"), encoding="utf-8").read()
    finally:
        try: hf.close()
        except Exception: pass
        try: os.unlink(tmp.name)
        except Exception: pass
    m = re.search(r"<body[^>]*>(.*)</body>", x, re.S); x = m.group(1) if m else x
    x = re.sub(r"<img[^>]*>", "", x)
    x = re.sub(r'\s(?:class|style|id|width|height|border|cellspacing|cellpadding|xmlns)="[^"]*"', "", x)
    x = re.sub(r"<(/?)(?:span|div|font)\b[^>]*>", "", x)
    x = re.sub(r"<p>\s*</p>", "", x)
    x = x.replace("<table>", "<table class='doc-table'>")
    return x.strip()

def hwpx_html(b):
    """HWPX(XML) → HTML: hp:tbl → <table>, hp:p → <p>"""
    import xml.etree.ElementTree as ET
    zz = zipfile.ZipFile(io.BytesIO(b)); out = []
    NS = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
    def ptext(p): return "".join(t.text or "" for t in p.iter(NS + "t"))
    def walk(node, out):
        for ch in list(node):
            if ch.tag == NS + "tbl":
                rows = []
                for tr in ch.iter(NS + "tr"):
                    cells = []
                    for tc in tr.findall(NS + "tc"):
                        cells.append("<br>".join(_esc(ptext(p)) for p in tc.iter(NS + "p")))
                    rows.append("<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>")
                out.append("<table class='doc-table'><tbody>" + "".join(rows) + "</tbody></table>")
            elif ch.tag == NS + "p":
                if ch.find(".//" + NS + "tbl") is not None:
                    walk(ch, out)
                else:
                    t = ptext(ch).strip()
                    if t: out.append(f"<p>{_esc(t)}</p>")
            else:
                walk(ch, out)
    for n in sorted(zz.namelist()):
        if n.startswith("Contents/section"):
            try:
                walk(ET.fromstring(zz.read(n)), out)
            except Exception as e:
                out.append(f"<p>{_esc(hwpx_text(zz.read(n).decode('utf-8','replace')))}</p>")
    return "".join(out)

def extract_html(name, b):
    ext = os.path.splitext(name)[1].lower()
    try:
        if ext == ".pdf": return pdf_html(b)
        if ext == ".hwp": return hwp_html(b)
        if ext == ".hwpx": return hwpx_html(b)
        if ext in (".xls", ".xlsx"):
            t = extract(name, b)
            rows = [r.split(chr(9)) for r in t.split(chr(10)) if r.strip()]
            return "<table class='doc-table'><tbody>" + "".join("<tr>" + "".join(f"<td>{_esc(c)}</td>" for c in r) + "</tr>" for r in rows) + "</tbody></table>"
    except Exception as e:
        print("   ! HTML 변환 실패", name, e)
    return ""

def clean_text(t):
    t = t.replace("\x00", "").replace("﻿", "")
    t = re.sub(r"<그림>|<표>|</?[a-z]+>", " ", t)
    t = re.sub(r"[ \t　]+", " ", t)
    t = re.sub(r"\n\s*\n+", "\n", t)
    return t.strip()

def clean_body(h):
    i = h.find('<div class="write_area">')
    if i < 0: return ""
    j = i + len('<div class="write_area">'); depth = 1
    for m in re.finditer(r"<div\b|</div>", h[j:]):
        depth += 1 if m.group(0).startswith("<div") else -1
        if depth == 0: body = h[j:j + m.start()]; break
    else:
        body = h[j:]
    body = re.sub(r"<script.*?</script>", "", body, flags=re.S | re.I)
    body = re.sub(r"<(?:link|meta|style)[^>]*>.*?(?:</style>)?", "", body, flags=re.S | re.I)
    body = re.sub(r'\s(?:style|class|onmouseover|onmouseout|width|height)="[^"]*"', "", body)
    # 본문 안 그림: 그룹웨어(cimg)에서 받아 둔 파일(files/gw/img/)이 있으면 그것으로 바꾸고, 없으면 지움
    def _img(m):
        src = re.search(r"src=[\"']([^\"']+)", m.group(0))
        if not src: return ""
        name = src.group(1).split("?")[0].rsplit("/", 1)[-1]
        if "/cimg/" in src.group(1) and os.path.exists(os.path.join(FILES, "img", name)):
            return f'<img src="../../files/gw/img/{name}" alt="" style="max-width:100%;height:auto;display:block;margin:10px 0">'
        return ""
    body = re.sub(r"<img[^>]*>", _img, body, flags=re.I)
    body = re.sub(r"<a [^>]*JAVASCRIPT[^>]*>(.*?)</a>", r"\1", body, flags=re.S)
    return body.strip()

def main():
    os.makedirs(FILES, exist_ok=True)
    out = []
    for zp in ZIPS:
      z = zipfile.ZipFile(zp)
      d = json.loads(z.read("posts.json").decode("utf-8"))
      posts = json.loads(d["posts"]) if isinstance(d["posts"], str) else d["posts"]
      for p in posts:
          if p.get("cat"):
              code, sub = p["cat"], {"budget": "예산결산서", "rules": "노사협의회 운영규약", "news": "단체교섭", "council": "공지(노사협의회)", "guide": p.get("boardName", "가이드라인"), "documents": p.get("boardName", "참고자료"), "statement": p.get("boardName", "언론동향")}.get(p["cat"], p["cat"])
          elif p["board"] in BOARD_MAP:
              code, sub = BOARD_MAP[p["board"]]
          else:
              continue
          # 제목 기준 재분류: 운영규약 → 규약·규정, 단체교섭·단체협약 → 노조소식
          if code == "council":
              if re.search(r"운영규약|노사협의회\s*규약", p["title"]): code, sub = "rules", "노사협의회 운영규약"
              elif re.search(r"단체교섭|단체협약|임금협약|임단협|교섭", p["title"]): code, sub = "news", "단체교섭"
          raw = z.read("raw/" + p["id"] + ".html").decode("utf-8", "replace")
          body_html = clean_body(raw)
          body_text = clean_text(html.unescape(re.sub(r"<[^>]+>", " ", body_html)))
          rec = {"id": p["id"], "code": code, "sub": sub, "title": p["title"], "author": p["author"], "dept": p.get("dept", ""),
                 "date": p["date"].replace(".", "-"), "body_html": body_html, "body_text": body_text, "atts": []}
          print(f"[{code}] {rec['date']} {p['title'][:40]}")
          for a in p.get("atts", []):
              if "file" not in a: continue
              b = z.read("files/" + a["file"])
              safe = re.sub(r'[\\/:*?"<>|]', "_", a["file"])
              too_big = len(b) > MAX_ATT
              if not too_big: open(os.path.join(FILES, safe), "wb").write(b)
              clip = "언론기사" in a["name"]                     # 언론 기사 스크랩(저작권): 파일만, 본문 전문은 싣지 않음
              text = "" if clip else clean_text(extract(a["name"], b))[:300000]
              ahtml = extract_html(a["name"], b) if text else ""
              if ahtml and len(ahtml) > 1500000: ahtml = ""
              rec["atts"].append({"name": a["name"], "file": ("" if too_big else "files/gw/" + safe), "size": len(b), "text": text, "html": ahtml, "big": too_big})
              print(f"    - {a['name']} ({len(b)//1024} KB) → {len(text)}자")
          out.append(rec)
    out.sort(key=lambda r: (r["date"], r["id"]), reverse=True)
    json.dump({"posts": out}, open(os.path.join(ROOT, "data", "gw_posts.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("총", len(out), "건 → data/gw_posts.json")

if __name__ == "__main__":
    main()
